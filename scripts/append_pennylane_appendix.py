#!/usr/bin/env python3
"""Appends Appendix B (PennyLane-rendered circuit diagrams + specs, built by
build_pennylane_circuit_diagrams.py) to QGNN_V4_Final_Research_Report.docx,
matching the existing report's heading/caption/table visual style.

Appended at the end (after Appendix A) with its own B-prefixed figure/table
numbering so none of the document's existing Figure 1-20 / Table 1-13
numbers or cross-references need to change.
"""
from __future__ import annotations

import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(ROOT, "QGNN_V4_Final_Research_Report.docx")
DIAGRAMS_DIR = os.path.join(ROOT, "report_assets", "diagrams")

HEADER_FILL = "2E5395"
BODY_FONT_SIZE = Pt(9.5)

with open(os.path.join(DIAGRAMS_DIR, "appendix_b_circuit_specs.json")) as f:
    SPECS = json.load(f)


def set_heading(paragraph, style_id):
    """The report's styles.xml has duplicate-name style entries (two
    'Heading 1' elements, two 'Heading 2' elements), which breaks
    python-docx's name-based style lookup (doc.add_heading / p.style = ...).
    Setting w:pStyle by styleId directly sidesteps that lookup entirely."""
    pPr = paragraph._p.get_or_add_pPr()
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), style_id)


def add_heading(doc, text, level):
    p = doc.add_paragraph(text)
    set_heading(p, f"Heading{level}")
    return p


def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def style_header_row(row):
    for cell in row.cells:
        set_cell_shading(cell, HEADER_FILL)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.5)


def add_table(doc, headers, rows, col_widths_in):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    table.autofit = False
    for i, w in enumerate(col_widths_in):
        table.columns[i].width = Inches(w)
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].width = Inches(col_widths_in[i])
    style_header_row(table.rows[0])
    for row_vals in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_vals):
            cells[i].text = str(val)
            cells[i].width = Inches(col_widths_in[i])
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = BODY_FONT_SIZE
    return table


def add_caption(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"{label}. ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    return p


def add_centered_image(doc, path, width_in=6.0):
    p = doc.add_paragraph()
    # This document's sectPr sets a fixed w:docGrid linePitch; paragraphs
    # without an explicit w:spacing/lineRule="auto" override get their line
    # height (and any inline image in it) clipped to that grid pitch by
    # LibreOffice's renderer -- reproduced and root-caused against this
    # exact file, see the report's own working figures for the same
    # explicit-spacing pattern. Without this, the image is embedded
    # correctly (verified independently with mammoth) but renders as a
    # near-invisible sliver.
    pPr = p._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "120")
    spacing.set(qn("w:before"), "120")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    return p


def gate_str(gate_types: dict) -> str:
    order = ["RY", "RZ", "Rot", "CNOT"]
    parts = [f"{g}×{gate_types[g]}" for g in order if g in gate_types]
    return ", ".join(parts)


def main():
    doc = Document(DOCX_PATH)

    # Cross-reference from the existing Section 7 limitation paragraph to
    # the new appendix, without touching any existing figure/table numbers.
    for p in doc.paragraphs:
        if p.text.startswith("Introduce: reading left to right, the circuit is encode once"):
            r = p.add_run(
                " See Appendix B for the same circuit rendered directly by "
                "PennyLane's own qml.draw_mpl from the production code "
                "(src/scm_dataset/modeling/quantum/circuit.py), which "
                "removes this hand-verification caveat entirely."
            )
            r.italic = True
            break

    doc.add_page_break()
    add_heading(doc, "Appendix B — PennyLane-Rendered Quantum Circuit Diagrams and Specifications", 1)

    doc.add_paragraph(
        "Every diagram and every number in this appendix is generated directly from this "
        "project's own production code — src/scm_dataset/modeling/quantum/circuit.py's "
        "build_quantum_layer() — using PennyLane's own qml.draw_mpl and qml.specs "
        "functions (scripts/build_pennylane_circuit_diagrams.py), not hand-drawn or manually "
        "transcribed. Gates are shown at PennyLane's “device” decomposition level, so "
        "every template (AngleEmbedding, StronglyEntanglingLayers, ...) is expanded into the "
        "literal elementary gates (RY, RZ, Rot, CNOT) executed on the simulator, for exactly the "
        "qubit/layer/ansatz configurations this project actually ran. Parameter values shown on "
        "each gate are from a fixed random seed (torch.manual_seed(0)) used only to make the "
        "diagram concrete — the gate structure, qubit count, and parameter count are what "
        "matter and are identical for every seed."
    )

    # --- B.1 Reference circuit ---------------------------------------
    add_heading(doc, "B.1 Reference Circuit: StronglyEntanglingLayers (6 Qubits, 2 Layers)", 2)
    ref = SPECS["reference_strongly_entangling"]
    doc.add_paragraph(
        "This is the exact circuit behind every headline Primary and Severity result in this "
        "report (configs/qgnn_v4.yaml and configs/qgnn_v4_severity.yaml: quantum_v4.n_qubits=6, "
        "n_layers=2, ansatz=strongly_entangling), decomposed and drawn by PennyLane itself rather "
        "than approximated in Figure 6."
    )
    add_centered_image(doc, os.path.join(DIAGRAMS_DIR, ref["fig_name"]), width_in=6.3)
    add_caption(
        doc,
        "Figure B1",
        "PennyLane-rendered (qml.draw_mpl, device-level decomposition) circuit for the standing "
        "QGNN-v4 reference head: RY angle embedding on 6 qubits, two variational layers of a "
        "general single-qubit rotation (Rot) followed by a ring of CNOT entangling gates, and "
        "Pauli-Z measurement on every qubit. This is the literal gate sequence PennyLane "
        "executes — not a logical approximation.",
    )
    doc.add_paragraph(
        "Table B1 lists the circuit's exact specifications, read directly from qml.specs() at "
        "device-decomposition level: every gate PennyLane counts, the resulting circuit depth, "
        "and the number of trainable quantum parameters registered as ordinary PyTorch "
        "parameters via qml.qnn.TorchLayer."
    )
    add_table(
        doc,
        ["Property", "Value"],
        [
            ["Qubits (wires)", ref["num_wires"]],
            ["Variational layers", ref["n_layers"]],
            ["Ansatz", "StronglyEntanglingLayers (PennyLane built-in template)"],
            ["Encoding", "AngleEmbedding, rotation=“Y” (1 RY gate per qubit)"],
            ["Gate composition (device level)", gate_str(ref["gate_types"])],
            ["Total gate count", ref["num_gates"]],
            ["Circuit depth", ref["depth"]],
            ["Trainable quantum parameters", ref["trainable_params"]],
            ["Measurement", "expval(PauliZ) on all 6 qubits"],
            ["Device", ref["device"]],
            ["Differentiation method", ref["diff_method"] + " (exact analytic gradient on the state-vector simulator)"],
        ],
        col_widths_in=[2.6, 3.7],
    )
    add_caption(doc, "Table B1", "Exact specifications of the reference QGNN-v4 quantum circuit, from qml.specs().")

    # --- B.2 Ablation ansatz variants ---------------------------------
    add_heading(doc, "B.2 Ablation-Tested Ansatz Variants (Section 13.6, Phase 4 Stage 5)", 2)
    doc.add_paragraph(
        "Section 13.6's ansatz/entanglement-topology ablation (QGNN_V4_ABLATION_TABLE.csv) "
        "compared the reference circuit above against two other ansatz choices, both already "
        "implemented in the same production module and both actually run for that ablation "
        "— rendered here the same way, from the same code, for direct visual comparison."
    )

    ring = SPECS["hardware_efficient_ring"]
    add_centered_image(doc, os.path.join(DIAGRAMS_DIR, ring["fig_name"]), width_in=6.5)
    add_caption(
        doc,
        "Figure B2",
        "PennyLane-rendered hardware_efficient_ring variant: RY then RZ per qubit per layer "
        "(2 trainable rotation parameters/qubit/layer) followed by a CNOT ring.",
    )

    red = SPECS["reduced_entanglement"]
    add_centered_image(doc, os.path.join(DIAGRAMS_DIR, red["fig_name"]), width_in=5.6)
    add_caption(
        doc,
        "Figure B3",
        "PennyLane-rendered reduced_entanglement variant: RY only per qubit per layer "
        "(1 trainable rotation parameter/qubit/layer) followed by a CNOT chain (no wrap-around).",
    )

    doc.add_paragraph(
        "Table B2 compares the three ansatz variants' exact gate composition and parameter "
        "counts side by side — the same 36/24/12 trainable-parameter figures Section 13.6 "
        "reports, here traced directly to the gate-level circuit that produces them."
    )
    add_table(
        doc,
        ["Ansatz", "Trainable params", "Gate composition (device level)", "Depth"],
        [
            ["strongly_entangling (reference)", ref["trainable_params"], gate_str(ref["gate_types"]), ref["depth"]],
            ["hardware_efficient_ring", ring["trainable_params"], gate_str(ring["gate_types"]), ring["depth"]],
            ["reduced_entanglement", red["trainable_params"], gate_str(red["gate_types"]), red["depth"]],
        ],
        col_widths_in=[2.1, 1.3, 2.4, 0.8],
    )
    add_caption(doc, "Table B2", "Gate-level comparison of the three ansatz variants tested in Section 13.6's ablation, all read from qml.specs() on the actual production circuit.")

    doc.add_paragraph(
        "Reproduction: python scripts/build_pennylane_circuit_diagrams.py regenerates every "
        "figure and the specs JSON in this appendix directly from "
        "src/scm_dataset/modeling/quantum/circuit.py."
    ).runs[0].italic = True

    doc.save(DOCX_PATH)
    print("Saved", DOCX_PATH)


if __name__ == "__main__":
    main()
